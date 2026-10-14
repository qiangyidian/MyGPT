"""幻灯片的形状坐标：让 chunk 能落到具体形状上。"""
from __future__ import annotations

from app.rag.locator import locate
from app.rag.parsers import default_parser


def _make_pptx(path) -> str:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    for i in range(2):
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = f"第{i}张幻灯片的标题内容"
        box2 = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1))
        box2.text_frame.text = f"第{i}张幻灯片的正文内容"
    prs.save(str(path))
    return str(path)


def test_pptx_layout_shapes_are_in_slide_bounds_and_ordered(tmp_path):
    parsed = default_parser.parse(_make_pptx(tmp_path / "a.pptx"), ".pptx")

    assert parsed.layout is not None
    assert len(parsed.layout.slides) == 2
    slide = parsed.layout.slides[0]
    assert slide.slide_no == 1
    assert len(slide.shapes) == 2, "两个文本框应当各是一个形状"

    for shape in slide.shapes:
        assert 0 <= shape.x < shape.x + shape.w <= slide.width
        assert 0 <= shape.y < shape.y + shape.h <= slide.height
        assert parsed.text[shape.char_start : shape.char_end] == shape.text


def test_pptx_layout_text_is_unchanged(tmp_path):
    parsed = default_parser.parse(_make_pptx(tmp_path / "b.pptx"), ".pptx")
    assert parsed.text.startswith("# 幻灯片 1")
    assert "第0张幻灯片的标题内容" in parsed.text
    assert "第1张幻灯片的正文内容" in parsed.text


def test_pptx_chunk_on_second_shape_highlights_only_that_shape(tmp_path):
    """端到端：命中正文那一行时，只标正文形状，标题形状不能被标上。"""
    parsed = default_parser.parse(_make_pptx(tmp_path / "c.pptx"), ".pptx")
    slide = parsed.layout.slides[0]
    body = slide.shapes[1]

    loc = locate(parsed, page=None, char_start=body.char_start, char_end=body.char_end)

    assert loc is not None and loc.kind == "slide"
    assert loc.slide_no == 1
    hits = [s for s in loc.shapes if s.hit]
    assert len(hits) == 1
    assert "正文" in hits[0].text
    assert any(not s.hit and "标题" in s.text for s in loc.shapes)


def test_pptx_notes_are_not_shapes(tmp_path):
    """备注行不是形状，不该出现在形状列表里（它没有自己的版面框）。"""
    parsed = default_parser.parse(_make_pptx(tmp_path / "d.pptx"), ".pptx")
    for slide in parsed.layout.slides:
        assert all("备注" not in s.text for s in slide.shapes)
